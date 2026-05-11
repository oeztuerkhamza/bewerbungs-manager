# -*- coding: utf-8 -*-
"""Alle E-Mails aus dem Posteingang → Excel-Übersicht"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT = r"C:\Users\hamza\Desktop\Lebenslauf\Bewerbungen_Uebersicht.xlsx"

# (Datum, Firma, Position, Ergebnis, Notiz)
data = [
    # ── MAIL 1 ──
    ("08.01.2026", "WEB.DE", "-", "Info", "Willkommensmail FreeMail Postfach"),
    # ── MAIL 2 ──
    ("11.01.2026", "Optimus Search (XING)", "Back-End Developer C#/.NET", "Bewerbung verschickt", "Über XING beworben"),
    # ── MAIL 3 ──
    ("11.01.2026", "Optimus Search Ltd.", "Back-End Developer C#/.NET", "Eingangsbestätigung", "Automatische Bestätigung via onlyfy"),
    # ── MAIL 4 ──
    ("11.01.2026", "onlyfy by XING", "-", "Info", "Talent-Hub Einladung nach Bewerbung"),
    # ── MAIL 5 ──
    ("11.01.2026", "Ratbacher GmbH", "IT-Vermittlung", "Eingangsbestätigung", "E-Mail-Verifizierung angefordert"),
    # ── MAIL 6 ──
    ("14.01.2026", "WEB.DE", "-", "Newsletter", "Mobilfunk-Tarif Werbung"),
    # ── MAIL 7 ──
    ("16.01.2026", "Loy & Hutz Solutions GmbH", "Fullstack-Entwickler .NET/TypeScript", "Eingangsbestätigung", "Passwort erstellen für softgarden"),
    # ── MAIL 8 ──
    ("16.01.2026", "Loy & Hutz Solutions GmbH", "Fullstack-Entwickler .NET/TypeScript", "Eingangsbestätigung", "Bewerbung eingegangen, per du"),
    # ── MAIL 9 ──
    ("16.01.2026", "IQVIA", "Junior Software Developer C#/JS/Blazor/SQL", "Eingangsbestätigung", "Workday-Portal"),
    # ── MAIL 10 ──
    ("18.01.2026", "Tempton Next Level Experts GmbH", "Initiativbewerbung", "Eingangsbestätigung", "Denis Li, Recruiting"),
    # ── MAIL 11 ──
    ("18.01.2026", "Stiegeler Internet Service GmbH", "Softwareentwickler", "Eingangsbestätigung", "Automatische Bestätigung"),
    # ── MAIL 12 ──
    ("19.01.2026", "Stiegeler Internet Service GmbH", "Softwareentwickler", "Unterlagen nachgefordert", "Anschreiben, Lebenslauf, Zeugnisse fehlen – Frist 22.01."),
    # ── MAIL 13 ──
    ("19.01.2026", "Tempton Next Level Experts GmbH", "Initiativbewerbung", "Datenschutz-Einwilligung", "Jonas Reindl, Datenschutz-Einwilligung nötig"),
    # ── MAIL 14 ──
    ("19.01.2026", "Tempton Next Level Experts GmbH", "Initiativbewerbung", "Weitere Angaben angefordert", "Verfügbarkeit, Gehalt, Region etc. gefragt"),
    # ── MAIL 15 ──
    ("19.01.2026", "HRworks GmbH", "Junior Softwareentwickler", "Eingangsbestätigung", ""),
    # ── MAIL 16 ──
    ("20.01.2026", "IQVIA", "Junior Software Developer C#/JS/Blazor/SQL", "❌ Absage", "Nach 4 Tagen abgelehnt"),
    # ── MAIL 17 ──
    ("20.01.2026", "Google", "-", "Info", "E-Mail-Adresse bestätigen"),
    # ── MAIL 18-20 ──
    ("20.01.2026", "Deutsche Rentenversicherung BW", "Bewerbung", "Verifizierung", "3× Verifizierungscode gesendet"),
    # ── MAIL 21 ──
    ("20.01.2026", "Deutsche Rentenversicherung BW", "Bewerbung", "Info", "Passwort-Anforderung"),
    # ── MAIL 22 ──
    ("21.01.2026", "HERMA GmbH", "SAP Developer / Application Manager", "Eingangsbestätigung", "SuccessFactors-Portal"),
    # ── MAIL 23 ──
    ("21.01.2026", "Google", "-", "Info", "Familiengruppe – fatih Hug beigetreten"),
    # ── MAIL 24 ──
    ("21.01.2026", "HRworks GmbH (HR WORKS)", "Junior Softwareentwickler", "❌ Absage", "Anne Burger – nicht in engere Auswahl"),
    # ── MAIL 25 ──
    ("21.01.2026", "WEB.DE", "-", "Newsletter", "Mobilfunk-Tarif Werbung"),
    # ── MAIL 26 ──
    ("22.01.2026", "Netto Marken-Discount", "(Junior) Full-Stack-Entwickler C#", "Eingangsbestätigung", ""),
    # ── MAIL 27 ──
    ("23.01.2026", "Loy & Hutz Solutions GmbH", "Fullstack-Entwickler .NET/TypeScript", "📞 Interview-Einladung", "Video-Call 26.01. um 16:00 Uhr, Christina Schweizer + Achim Wüst"),
    # ── MAIL 28 ──
    ("23.01.2026", "Loy & Hutz Solutions GmbH", "-", "Info", "Passwort zurücksetzen softgarden"),
    # ── MAIL 29 ──
    ("23.01.2026", "Microsoft", "-", "Info", "E-Mail-Adresse bestätigen, Code 604838"),
    # ── MAIL 30 ──
    ("23.01.2026", "Microsoft", "-", "Info", "Marketing-Einstellungen bestätigen"),
    # ── MAIL 31 ──
    ("23.01.2026", "Microsoft", "-", "Info", "Willkommen Microsoft-Konto"),
    # ── MAIL 32 ──
    ("23.01.2026", "Loy & Hutz Solutions GmbH", "Fullstack-Entwickler .NET/TypeScript", "📞 Interview bestätigt", "Teams-Link nachgesendet"),
    # ── MAIL 33 ──
    ("26.01.2026", "Stiegeler Internet Service GmbH", "Softwareentwickler", "❌ Absage", "Unterlagen nicht nachgereicht → nicht berücksichtigt"),
    # ── MAIL 34 ──
    ("26.01.2026", "Microsoft", "-", "Info", "Einmalcode 366486"),
    # ── MAIL 35 ──
    ("26.01.2026", "Loy & Hutz Solutions GmbH", "Fullstack-Entwickler .NET/TypeScript", "Unterlagen nachgefordert", "Arbeitszeugnis nachreichen"),
    # ── MAIL 36 ──
    ("26.01.2026", "Bundesverwaltungsamt (BVA)", "AWV-2025-133", "Registrierung", "Go4Bund Passwort erhalten"),
    # ── MAIL 37 ──
    ("27.01.2026", "Loy & Hutz Solutions GmbH", "Fullstack-Entwickler .NET/TypeScript", "Feedback erbeten", "Bewerbungserlebnis bewerten"),
    # ── MAIL 38 ──
    ("28.01.2026", "WEB.DE", "-", "Newsletter", "Gewinnspiel + Sicherheits-Patches"),
    # ── MAIL 39 ──
    ("29.01.2026", "Tempton Next Level Experts GmbH", "Initiativbewerbung", "Erinnerung", "Erneut weitere Angaben angefordert"),
    # ── MAIL 40 ──
    ("29.01.2026", "Netto Marken-Discount", "(Junior) Full-Stack-Entwickler C#", "❌ Absage", "Vorauswahl – nicht berücksichtigt"),
    # ── MAIL 41 ──
    ("01.02.2026", "Google", "-", "Info", "Einstellungen aktualisiert (Alter nicht bestätigt)"),
    # ── MAIL 42 ──
    ("02.02.2026", "Microsoft OneDrive", "-", "Info", "Dateien mit OneDrive sicher"),
    # ── MAIL 43 ──
    ("02.02.2026", "CSB-System SE", "Bewerbung", "Eingangsbestätigung", "Jacqueline Mingers, Recruiting"),
    # ── MAIL 44 ──
    ("02.02.2026", "COPA Systeme GmbH & Co. KG", "Softwareentwickler", "Eingangsbestätigung", "Lebenslauf + Anschreiben nachgefordert"),
    # ── MAIL 45 ──
    ("02.02.2026", "RIB Cosinus GmbH", "Junior Anwendungsentwickler MS Dynamics 365 BC", "Eingangsbestätigung", "Online-Formular bestätigt"),
    # ── MAIL 46 ──
    ("03.02.2026", "HeyJobs", "-", "Registrierung", "Passwort erstellen"),
    # ── MAIL 47 ──
    ("03.02.2026", "HeyJobs / Venios GmbH", "Angular Frontend Developer", "Bewerbung verschickt", "11 Stunden zum Fragen beantworten"),
    # ── MAIL 48 ──
    ("03.02.2026", "Flughafen Stuttgart GmbH", "Sachbearbeiter IT-Servicecenter", "Eingangsbestätigung", "Sebastian Zeller, HR Crew"),
    # ── MAIL 49 ──
    ("03.02.2026", "HERMA GmbH", "SAP Developer / Application Manager", "❌ Absage", "Nicht in engere Auswahl"),
    # ── MAIL 50 ──
    ("03.02.2026", "GermanTechJobs", "Java Jobs", "Newsletter-Bestätigung", "Newsletter-Abo bestätigen"),
    # ── MAIL 51 ──
    ("03.02.2026", "Captana / VusionGroup", "Softwareentwickler Fullstack", "Eingangsbestätigung", "SmartRecruiters"),
    # ── MAIL 52 ──
    ("03.02.2026", "TGW Logistics", "Software Developer Logistiksoftware (WMS/MFR)", "Eingangsbestätigung", "rexx-systems"),
    # ── MAIL 53 ──
    ("03.02.2026", "HRworks GmbH", "Junior Softwareentwickler", "Eingangsbestätigung", "2. Bewerbung"),
    # ── MAIL 54 ──
    ("03.02.2026", "Loy & Hutz Solutions GmbH", "Fullstack-Entwickler .NET/TypeScript", "❌ Absage", "Nach Gespräch – andere Profile besser passend"),
    # ── MAIL 55 ──
    ("03.02.2026", "Stadt Freiburg (wirliebenfreiburg.de)", "Initiativ / Fachinformatiker", "Kontaktformular", "Bestätigung Kontaktformular"),
    # ── MAIL 56 ──
    ("03.02.2026", "TIGNUM GmbH", "Full Stack Developer", "❌ Absage", "Anderer Kandidat gewählt"),
    # ── MAIL 57 ──
    ("22.01.2026", "Bartels-Langness (bela.de)", "Web-Developer", "❌ Absage", "Recruiting, hohe Bewerbungsqualität"),
    # ── MAIL 58 ──
    ("04.02.2026", "TGW Logistics", "Software Developer Logistiksoftware (WMS/MFR)", "❌ Absage", "Nicht in engere Wahl"),
    # ── MAIL 59 ──
    ("04.02.2026", "HeyJobs / Venios GmbH", "Angular Frontend Developer", "❌ Absage", "Nicht in engerer Auswahl"),
    # ── MAIL 60 ──
    ("04.02.2026", "Stadt Freiburg", "Fachinformatiker (Initiativ)", "Antwort", "Aktuell keine Stelle, Jobalert empfohlen, Initiativbewerbung möglich"),
    # ── MAIL 61 ──
    ("04.02.2026", "Stadt Freiburg i.Br.", "Initiativbewerbung", "Eingangsbestätigung", "rexx-systems"),
    # ── MAIL 62 ──
    ("04.02.2026", "Markant Gruppe", "Software Engineer", "Bestätigung angefordert", "Bewerbung innerhalb 20 Tagen bestätigen"),
    # ── MAIL 63 ──
    ("05.02.2026", "seidemann: solutions GmbH", "(Junior) ERP Software Developer", "❌ Absage", "Keine passende Stelle"),
    # ── MAIL 64 ──
    ("05.02.2026", "Tempton Next Level Experts GmbH", "Initiativbewerbung", "Letzte Erinnerung", "Erneut Angaben angefordert"),
    # ── MAIL 65 ──
    ("05.02.2026", "Instaffo", "-", "Info", "E-Mail-Adresse ändern bestätigen"),
    # ── MAIL 66 ──
    ("05.02.2026", "Tempton Next Level Experts GmbH", "Initiativbewerbung", "Im Bewerberpool", "Fragen beantwortet, Arbeitserlaubnis + Zertifikate nachgefordert"),
    # ── MAIL 67 ──
    ("05.02.2026", "HRworks GmbH (HR WORKS)", "Junior Softwareentwickler", "❌ Absage", "2. Bewerbung auch abgelehnt"),
    # ── MAIL 68 ──
    ("06.02.2026", "TOP Mehrwert Logistik", "Kennenlernen", "📞 Interview-Einladung", "Video 11.02. um 13:00, Annalisa Richter"),
    # ── MAIL 69 ──
    ("06.02.2026", "WEB.DE", "-", "Newsletter", "Mobilfunk + Zattoo TV"),
    # ── MAIL 70 ──
    ("06.02.2026", "eSIM.sm", "-", "Registrierung", "Willkommen bei eSIM.sm"),
    # ── MAIL 71 ──
    ("06.02.2026", "eSIM.sm", "-", "Kaufbestätigung", "eSIM #1372809 Switzerland Unlimited 1 day"),
    # ── MAIL 72 ──
    ("06.02.2026", "eSIM.sm (Stripe)", "-", "Rechnung", "4,40 € via Apple Pay"),
    # ── MAIL 73 ──
    ("07.02.2026", "eSIM.sm", "-", "Info", "eSIM erfolgreich installiert"),
    # ── MAIL 74 ──
    ("07.02.2026", "eSIM.sm", "-", "Info", "eSIM aktiviert – Schweiz"),
    # ── MAIL 75 ──
    ("08.02.2026", "Lekker Code Company (XING)", "Fullstack Developer C#/Angular", "Bewerbung verschickt", "Über XING beworben"),
    # ── MAIL 76 ──
    ("09.02.2026", "COPA Systeme GmbH & Co. KG", "Softwareentwickler", "❌ Absage", "Kein passender Funktionsbereich"),
    # ── MAIL 77 ──
    ("09.02.2026", "Instaffo / LapID Service GmbH", "Softwareentwickler .NET/C#", "Antwort ausstehend", "Unternehmen hat geschrieben, Antwort fehlt"),
    # ── MAIL 78 ──
    ("09.02.2026", "DIS AG", "IT Administrator - ERP", "❌ Absage", "Im Bewerberpool aufgenommen"),
    # ── MAIL 79 ──
    ("09.02.2026", "AllatNet Recruiting GmbH", "-", "Datenschutz-Bestätigung", "Datenschutzrichtlinie akzeptieren"),
    # ── MAIL 80 ──
    ("09.02.2026", "Flughafen Stuttgart GmbH", "Sachbearbeiter IT-Servicecenter", "❌ Absage", "Nicht berücksichtigt, Job-Newsletter empfohlen"),
    # ── MAIL 81 ──
    ("10.02.2026", "Wavestone (Q-Perior)", "Full-Stack-Entwickler .NET/React", "❌ Absage", "Anderer Projektpartner gewählt"),
    # ── MAIL 82 ──
    ("11.02.2026", "DHL Paket", "-", "Zustellung", "Sendung am Ablageort (Briefkasten)"),
    # ── MAIL 83 ──
    ("11.02.2026", "Instaffo / XignSys GmbH", "Mobile App Experte (Android & iOS)", "Job-Vorschlag", "45.000–62.000 €, Hybrid"),
    # ── MAIL 84 ──
    ("12.02.2026", "Markant Gruppe", "Software Engineer", "Erinnerung", "Bewerbung bestätigen!"),
    # ── MAIL 85 ──
    ("12.02.2026", "Markant Gruppe", "Software Engineer", "Eingangsbestätigung", "Bewerbung bestätigt + angenommen"),
    # ── MAIL 86 ──
    ("12.02.2026", "Markant Gruppe", "-", "Registrierung", "Karriereportal Passwort erhalten"),
    # ── MAIL 87 ──
    ("13.02.2026", "Instaffo / Datora WebSystems", "Shopify Frontend Entwickler", "Job-Vorschlag", "40.000–60.000 €, Remote"),
    # ── MAIL 88 ──
    ("13.02.2026", "Instaffo / Datora WebSystems", "Shopify Frontend Entwickler", "Antwort ausstehend", "Unternehmen hat geschrieben"),
    # ── MAIL 89 ──
    ("14.02.2026", "Instaffo / Datora WebSystems", "Shopify Frontend Entwickler", "Erinnerung", "Nachricht von Pia Loddenkemper unbeantwortet"),
    # ── MAIL 90 ──
    ("16.02.2026", "TOP Mehrwert Logistik", "Kennenlernen (2. Runde)", "📞 2. Interview-Einladung", "Video 19.02. um 13:00, + techn. Leiter Günther Engelhardt"),
    # ── MAIL 91 ──
    ("18.02.2026", "Stadt Freiburg i.Br.", "Initiativbewerbung", "❌ Absage", "Unbefristete Stellen müssen ausgeschrieben werden"),
    # ── MAIL 92 ──
    ("18.02.2026", "Bundesagentur für Arbeit (IT-Systemhaus)", "IT-Systemhaus Bewerbung", "❌ Absage", "Martina Beyer, Fachkraft Rekrutierung"),
    # ── MAIL 93 ──
    ("19.02.2026", "WEB.DE", "-", "Newsletter", "Traumhaus-Verlosung Föhr"),
    # ── MAIL 94 ──
    ("20.02.2026", "Markant Gruppe", "Software Engineer", "❌ Absage", "Andere Kandidaten passgenauer"),
    # ── MAIL 95 ──
    ("23.02.2026", "CHECK24", "(Junior) Web Entwickler", "Eingangsbestätigung", ""),
    # ── MAIL 96 ──
    ("23.02.2026", "onOffice GmbH", "Bewerbung", "Eingangsbestätigung", "Fachteam prüft Unterlagen"),
    # ── MAIL 97 ──
    ("23.02.2026", "DEKRA", "Junior Full Stack Entwickler", "Eingangsbestätigung", "Job ID 4167"),
    # ── MAIL 98 ──
    ("23.02.2026", "Vimata Group / Step Ahead", "ERP Consulting", "📞 Interview-Einladung", "Teams-Link, Philip Koprek"),
    # ── MAIL 99 ──
    ("23.02.2026", "DSV-Gruppe", "Softwareentwickler", "Eingangsbestätigung", "Workday, Irina Felde"),
    # ── MAIL 100 ──
    ("23.02.2026", "DSV-Gruppe", "Softwareentwickler", "Eingangsbestätigung", "Bestätigungsmail"),
    # ── MAIL 101 ──
    ("23.02.2026", "Microsoft", "-", "Info", "Einmalcode 003880"),
    # ── MAIL 102 ──
    ("24.02.2026", "CONET", "(Junior) Frontend Entwickler JavaScript", "Eingangsbestätigung", ""),
    # ── MAIL 103 ──
    ("24.02.2026", "CONET", "(Junior) Full Stack Entwickler", "Eingangsbestätigung", ""),
    # ── MAIL 104 ──
    ("24.02.2026", "Bundesagentur für Arbeit", "-", "Info", "E-Mail-Adresse bestätigen"),
    # ── MAIL 105 ──
    ("24.02.2026", "inovex GmbH", "Software Developer Full-Stack (intern)", "Eingangsbestätigung", "Recruitee"),
    # ── MAIL 106 ──
    ("24.02.2026", "DIRINGER & SCHEIDEL IT Services", "Full-Stack Developer Vue.js/.NET", "Eingangsbestätigung", "onlyfy Bewerbungsmanager"),
    # ── MAIL 107 ──
    ("24.02.2026", "DIRINGER & SCHEIDEL IT Services", "Full-Stack Developer Vue.js/.NET", "Eingangsbestätigung", "Bestätigungsmail"),
    # ── MAIL 108 ──
    ("24.02.2026", "EBZ Gruppe", "Fullstack Entwickler", "Eingangsbestätigung", "SuccessFactors, Jana Schmid"),
    # ── MAIL 109 ──
    ("24.02.2026", "EBZ Gruppe", "Fullstack Entwickler", "Eingangsbestätigung", "Bewerbung vom 24.02."),
    # ── MAIL 110 ──
    ("24.02.2026", "Sanero Medical GmbH", "Softwareentwickler med. Apps (DiGA)", "❌ Absage", "Andy Bosch, sofortige Absage"),
    # ── MAIL 111 ──
    ("24.02.2026", "Roxtra GmbH", "Software Developer KI-Integration/Schnittstellen", "Eingangsbestätigung", "Inka Bezold, Personalmarketing"),
    # ── MAIL 112 ──
    ("24.02.2026", "GermanTechJobs", "-", "Erinnerung", "Newsletter-Abo noch gewünscht?"),
    # ── MAIL 113 ──
    ("24.02.2026", "Instaffo / MACH AG", "Senior Java Full-Stack Entwickler", "Antwort ausstehend", "Patricia wartet auf Antwort"),
    # ── MAIL 114 ──
    ("25.02.2026", "WEB.DE", "-", "Newsletter", "Aktion Mensch + KI Fear Speech"),
    # ── MAIL 115 ──
    ("25.02.2026", "CHECK24", "(Junior) Web Entwickler", "❌ Absage", "Anna Scharfenstein, Recruiting"),
    # ── MAIL 116 ──
    ("25.02.2026", "ABOSCO GmbH", "Softwareentwickler Java/C# – Junior", "❌ Absage", "Otto Schmidlin"),
    # ── MAIL 117 ──
    ("26.02.2026", "Ariel (Procter & Gamble)", "-", "Rückerstattung", "Ariel MaxPower 9,95 € Rückerstattung"),
    # ── MAIL 118 ──
    ("26.02.2026", "Lenor (Procter & Gamble)", "-", "Rückerstattung", "Lenor Wäscheparfüm 4,25 € Rückerstattung"),
    # ── MAIL 119 ──
    ("27.02.2026", "inovex GmbH", "Software Developer Full-Stack", "❌ Absage", "Kenntnisse passen nicht vollständig"),
    # ── MAIL 120 ──
    ("27.02.2026", "Roxtra GmbH", "Software Developer React/TypeScript/C#/.NET/KI", "❌ Absage", "Keine Position anbietbar"),
    # ── MAIL 121 ──
    ("01.03.2026", "Lukrativ GmbH", "Full-Stack Entwickler", "❌ Absage", "Matthias Luchner"),
    # ── MAIL 122 ──
    ("01.03.2026", "DSV-Gruppe", "Softwareentwickler", "❌ Absage", "Anderer Bewerberkreis bevorzugt"),
    # ── MAIL 123 ──
    ("02.03.2026", "onOffice GmbH", "Bewerbung", "❌ Absage", "Sebastian Koch, HR"),
    # ── MAIL 124 ──
    ("02.03.2026", "CONET", "(Junior) Full Stack Entwickler", "❌ Absage", ""),
    # ── MAIL 125 ──
    ("02.03.2026", "DEKRA", "Junior Full Stack Entwickler", "❌ Absage", "Annalena Buca, Job ID 4167"),
    # ── MAIL 126 ──
    ("02.03.2026", "Vimata Group / Step Ahead", "ERP Consulting", "❌ Absage", "Philip Koprek – Fachbereich entschied dagegen, Bewerberpool angeboten"),
    # ── MAIL 127 ──
    ("03.03.2026", "CONET", "(Junior) Frontend Entwickler JavaScript", "❌ Absage", ""),
    # ── MAIL 128 ──
    ("03.03.2026", "Dicom GmbH (Stephanie Gierelt)", "ehem. Arbeitgeber", "Info", "Arbeitsbescheinigung + Zeugnis wird bearbeitet"),
    # ── MAIL 129 ──
    ("04.03.2026", "Instaffo / Atruvia AG", "Software Engineer Frontend Homepage", "Job-Vorschlag", "62.000–91.000 €, Hybrid, Aschheim"),
    # ── MAIL 130 ──
    ("05.03.2026", "EBZ Gruppe", "Fullstack Entwickler", "❌ Absage", "Andere Bewerber passender"),
    # ── MAIL 131 ──
    ("06.03.2026", "Instaffo / MACH AG", "Senior Java Full-Stack Entwickler", "❌ Absage (Prozess beendet)", "174 neue Job-Vorschläge"),
    # ── MAIL 132 ──
    ("06.03.2026", "GitHub", "-", "Sicherheit", "Device Verification, Code 911073"),
    # ── MAIL 133 ──
    ("07.03.2026", "Amazon.de", "-", "Bestellung", "Bremsbeläge Fahrrad bestellt, 10,99 €"),
    # ── MAIL 134 ──
    ("07.03.2026", "Amazon Prime", "-", "Abo", "Prime Shipping Studenten aktiviert, 4,49 €/Monat"),
    # ── MAIL 135 ──
    ("07.03.2026", "Netlify", "-", "Info", "Free Team auf Netlify erstellt"),
    # ── MAIL 136 ──
    ("08.03.2026", "Amazon Prime", "-", "Kündigung", "Prime gekündigt, 4,49 € erstattet"),
    # ── MAIL 137 ──
    ("08.03.2026", "Amazon.de", "-", "Versand", "Bremsbeläge versendet"),
    # ── MAIL 138 ──
    ("08.03.2026", "eSIM.sm", "-", "Marketing", "Rabattangebote für Reisen"),
    # ── MAIL 139 ──
    ("09.03.2026", "GitGuardian", "-", "Sicherheitswarnung", "SMTP credentials exposed – KulturPlatform"),
    # ── MAIL 140 ──
    ("09.03.2026", "GitGuardian", "-", "Sicherheitswarnung", "SMTP credentials exposed (2. Meldung)"),
    # ── MAIL 141 ──
    ("08.03.2026", "GitHub", "-", "Info", "GitGuardian App autorisiert"),
    # ── MAIL 142 ──
    ("09.03.2026", "GitGuardian", "-", "Sicherheitswarnung", "Historical Scan – 4 Secrets gefunden"),
    # ── MAIL 143 ──
    ("09.03.2026", "GitGuardian", "-", "Sicherheitswarnung", "Historical Scan – 5 Secrets gefunden"),
    # ── MAIL 144 ──
    ("09.03.2026", "GitGuardian (Dwayne McDaniel)", "-", "Info", "ggshield Shift Left Protection empfohlen"),
    # ── MAIL 145 ──
    ("09.03.2026", "Instaffo", "Diverse Jobs", "Job-Vorschläge", "Angular Developer, Frontend, Backend, Fullstack"),
    # ── MAIL 146 ──
    ("09.03.2026", "Amazon.de", "-", "Zustellung", "Bremsbeläge in Zustellung"),
    # ── MAIL 147 ──
    ("09.03.2026", "Amazon.de (Luna)", "-", "Kündigung", "Luna Premium-Abo gekündigt (Zahlungsproblem)"),
    # ── MAIL 148 ──
    ("09.03.2026", "Cloudflare", "-", "Aktion erforderlich", "Nameserver für melike-ve-musa-evleniyor.com aktualisieren"),
    # ── MAIL 149 ──
    ("09.03.2026", "Amazon.de", "-", "Zustellung", "Bremsbeläge geliefert, im Briefkasten"),
    # ── MAIL 150 ──
    ("10.03.2026", "Amazon.de", "-", "Marketing", "Frühlingsangebote"),
    # ── MAIL 151 ──
    ("10.03.2026", "Christian Funk Holding GmbH & Co KG", "Bewerbung", "Eingangsbestätigung", "HR Recruiting Team, Offenburg"),
    # ── MAIL 152 ──
    ("10.03.2026", "Reply Deutschland SE", "Junior Software Developer – Logistics Software", "Eingangsbestätigung", "job@reply.de"),
    # ── MAIL 153 ──
    ("10.03.2026", "FERCHAU GmbH (Karlsruhe IT)", "IT-Karriere", "Eingangsbestätigung", ""),
    # ── MAIL 154 ──
    ("10.03.2026", "Empfehlungsbund / GISA GmbH", "(Junior-) Entwickler:in", "Bewerbung gestartet", "Link zum Bewerbermanagementsystem"),
    # ── MAIL 155 ──
    ("10.03.2026", "G&S IT Group GmbH", "Bewerbung", "Eingangsbestätigung", "Recruiting-Team prüft"),
    # ── MAIL 156 ──
    ("10.03.2026", "msg systems ag", "Bewerbung 1", "Datenschutz angefordert", "Zustimmung nötig für Bearbeitung"),
    # ── MAIL 157 ──
    ("10.03.2026", "msg systems ag", "Bewerbung 2", "Datenschutz angefordert", "Zustimmung nötig für Bearbeitung"),
    # ── MAIL 158 ──
    ("10.03.2026", "msg systems ag", "Bewerbung 3", "Datenschutz angefordert", "Zustimmung nötig für Bearbeitung"),
    # ── MAIL 159 ──
    ("11.03.2026", "Bundesagentur für Arbeit", "-", "Jobalert", "2 neue Treffer: ANG GmbH Software-Entwickler Vue.js/Go"),
    # ── MAIL 160 ──
    ("11.03.2026", "Instaffo / Atruvia AG", "Software Engineer Fullstack Online Banking", "Job-Vorschlag", "Über Wunschgehalt, Hybrid"),
    # ── MAIL 161 ──
    ("11.03.2026", "DIRINGER & SCHEIDEL IT Services", "Full-Stack Developer Vue.js/.NET", "❌ Absage", "Vivien Oelschlegel"),
    # ── MAIL 162 ──
    ("11.03.2026", "GitGuardian (Dwayne McDaniel)", "-", "Info", "Start detecting secrets – Ressourcen"),
    # ── MAIL 163 ──
    ("11.03.2026", "WEB.DE", "-", "Newsletter", "Traumhaus Föhr + VW California + Busuu"),
    # ── MAIL 164 ──
    ("11.03.2026", "Amazon.de", "-", "Marketing", "Amazon Haul Werbung"),
    # ── MAIL 165 ──
    ("11.03.2026", "HELLMUT RUCK GmbH", "Software Developer (Web & Data Integration)", "Eingangsbestätigung", "People & Culture"),
    # ── MAIL 166 ──
    ("11.03.2026", "abasoft EDV-Programme GmbH", "Fullstack-Entwickler AI-Driven Development", "Eingangsbestätigung", "Automatisch generiert"),
    # ── MAIL 167 ──
    ("12.03.2026", "Bundeskriminalamt (BKA)", "T-2026-5", "Registrierung", "Account aktivieren, 48h gültig"),
    # ── MAIL 168 ──
    ("12.03.2026", "Bundeskriminalamt (BKA)", "T-2026-5", "Eingangsbestätigung", "Kennung 133001"),
    # ── MAIL 169 ──
    ("12.03.2026", "MZA Meyer-Zweiradtechnik GmbH", "Junior Software-Entwickler .NET/C#", "Eingangsbestätigung", "rexx-systems"),
    # ── MAIL 170 ──
    ("12.03.2026", "finest jobs (rexx)", "-", "Registrierung", "Passwort setzen für finest-jobs Zugang"),
    # ── MAIL 171 ──
    ("12.03.2026", "Loy & Hutz Solutions GmbH", "Neue Bewerbung", "Info", "Passwort zurücksetzen softgarden"),
    # ── MAIL 172 ──
    ("12.03.2026", "Loy & Hutz Solutions GmbH", "Neue Bewerbung", "Eingangsbestätigung", "2. Bewerbung eingegangen"),
    # ── MAIL 173 ──
    ("12.03.2026", "Häfele SE & Co KG", "Developer – Java Full Stack", "Eingangsbestätigung", "Lara Elsässer, 2 Tage mobiles Arbeiten"),
    # ── MAIL 174 ──
    ("12.03.2026", "abasoft EDV-Programme GmbH", "Fullstack-Entwickler AI-Driven Development", "❌ Absage", "Florian Fiedler, Leiter Softwareentwicklung"),
    # ── MAIL 175 ──
    ("12.03.2026", "Häfele Technology Solutions", "Software Developer Testing", "Eingangsbestätigung", "Lena-Marie Bürkle, HR Business Partner"),
    # ── MAIL 176 ──
    ("13.03.2026", "Instaffo / mobisys GmbH", "Fullstack-Entwickler Mobile", "Antwort ausstehend", "Sebastian wartet auf Antwort"),
    # ── MAIL 177 ──
    ("13.03.2026", "AEB SE", "Bewerbung", "Eingangsbestätigung", "softgarden, Passwort erstellen"),
    # ── MAIL 178 ──
    ("13.03.2026", "AEB SE", "Bewerbung", "Eingangsbestätigung", "Bewerbung angekommen, barrierearm"),
    # ── MAIL 179 ──
    ("13.03.2026", "Interflex / Allegion", "KI Entwickler (all genders)", "Eingangsbestätigung", "Workday, Interflex HR-Team"),
    # ── MAIL 180 ──
    ("13.03.2026", "clean-tek Reinraumtechnik GmbH", "Fachinformatiker Anwendungsentwicklung", "Eingangsbestätigung", "Ann-Kathrin Gysau, HR"),
    # ── MAIL 181-183 ──
    ("13.03.2026", "msg systems ag", "3 Bewerbungen", "Datenschutz-Erinnerung", "7 Tage Frist für Datenschutz-Zustimmung"),
    # ── MAIL 184 ──
    ("13.03.2026", "Action1", "Full-Stack Developer Node.js/JS/TS", "Eingangsbestätigung", "Workable"),
    # ── MAIL 185 ──
    ("13.03.2026", "G&S IT Group GmbH", "Bewerbung", "❌ Absage", "Lurin, Career"),
    # ── MAIL 186 ──
    ("13.03.2026", "Action1", "Full-Stack Developer Node.js/JS/TS", "Eingangsbestätigung", "Team reviewing"),
    # ── MAIL 187 ──
    ("14.03.2026", "Administration Intelligence AG", "Fachinformatiker Anwendungsentwicklung", "Eingangsbestätigung", "Würzburg"),
    # ── MAIL 188 ──
    ("14.03.2026", "HUK-COBURG", "Frontend Entwickler Angular", "Verifizierung", "Identität bestätigen, Code 169090"),
]

# ── Excel erstellen ──────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Mail-Übersicht"

# Header
headers = ["#", "Datum", "Firma", "Position", "Ergebnis", "Notiz"]
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)

for col, h in enumerate(headers, 1):
    c = ws.cell(row=1, column=col, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = thin_border

# Farb-Zuordnung für Ergebnis
fill_absage  = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
fill_eingang = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
fill_interview = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
fill_warn    = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

def ergebnis_fill(val):
    v = val.lower()
    if "absage" in v:
        return fill_absage
    if "interview" in v:
        return fill_interview
    if "eingang" in v or "bestätigung" in v or "bewerbung verschickt" in v:
        return fill_eingang
    if "erinnerung" in v or "angefordert" in v or "ausstehend" in v or "nachgefordert" in v:
        return fill_warn
    return None

for i, (datum, firma, pos, ergebnis, notiz) in enumerate(data, 1):
    row = i + 1
    ws.cell(row=row, column=1, value=i).border = thin_border
    ws.cell(row=row, column=1).alignment = Alignment(horizontal="center")
    ws.cell(row=row, column=2, value=datum).border = thin_border
    ws.cell(row=row, column=3, value=firma).border = thin_border
    ws.cell(row=row, column=4, value=pos).border = thin_border

    ec = ws.cell(row=row, column=5, value=ergebnis)
    ec.border = thin_border
    ec.font = Font(name="Calibri", bold=True, size=10)
    f = ergebnis_fill(ergebnis)
    if f:
        ec.fill = f

    ws.cell(row=row, column=6, value=notiz).border = thin_border

# Spaltenbreiten
ws.column_dimensions["A"].width = 5
ws.column_dimensions["B"].width = 13
ws.column_dimensions["C"].width = 38
ws.column_dimensions["D"].width = 45
ws.column_dimensions["E"].width = 28
ws.column_dimensions["F"].width = 60

# Autofilter
ws.auto_filter.ref = f"A1:F{len(data)+1}"

# Zeilen einfrieren
ws.freeze_panes = "A2"

wb.save(OUTPUT)
print(f"✅ Excel gespeichert: {OUTPUT}")
print(f"   {len(data)} Einträge")
