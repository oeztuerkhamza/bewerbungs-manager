# -*- coding: utf-8 -*-
"""CSV'deki verileri mevcut Excel listesine ekle (mail listesinde olmayanları)"""

import csv, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

EXCEL = r"C:\Users\hamza\Desktop\Lebenslauf\Bewerbungen_Uebersicht.xlsx"
CSV_F = r"C:\Users\hamza\Desktop\Lebenslauf\Bewerbungen.csv"

# Durum çevirisi
DURUM_MAP = {
    "Olumsuz (Red)": "❌ Absage",
    "Alındı Teyidi (İşlemde)": "Eingangsbestätigung",
    "Başvuruldu": "Bewerbung verschickt",
    "Gönderildi": "Bewerbung verschickt",
    "Beworben": "Bewerbung verschickt",
    "Mülakat Daveti": "📞 Interview-Einladung",
}

# 1) Mevcut Excel'i oku
wb = openpyxl.load_workbook(EXCEL)
ws = wb.active

existing = set()
last_row = ws.max_row
for row in range(2, last_row + 1):
    firma = (ws.cell(row=row, column=3).value or "").strip().lower()
    pos = (ws.cell(row=row, column=4).value or "").strip().lower()
    existing.add((firma, pos))

# 2) CSV oku ve eksikleri bul
new_entries = []
with open(CSV_F, encoding="utf-8-sig") as f:
    reader = csv.DictReader(f)
    for row in reader:
        firma = row["Şirket / Kurum"].strip()
        pos = row["Pozisyon / Konu"].strip()
        durum = row["Durum / Sonuç"].strip()
        tarih = row["Tarih"].strip()
        
        key = (firma.lower(), pos.lower())
        if key not in existing:
            ergebnis = DURUM_MAP.get(durum, durum)
            new_entries.append((tarih, firma, pos, ergebnis, "CSV-Quelle"))
            existing.add(key)

# 3) Yeni satırları ekle
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
fill_absage = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
fill_eingang = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
fill_interview = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
fill_warn = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

def ergebnis_fill(val):
    v = val.lower()
    if "absage" in v: return fill_absage
    if "interview" in v: return fill_interview
    if "eingang" in v or "bestätigung" in v or "bewerbung" in v: return fill_eingang
    if "erinnerung" in v or "angefordert" in v or "ausstehend" in v: return fill_warn
    return None

num = last_row  # current max row number (including header)
for tarih, firma, pos, ergebnis, notiz in new_entries:
    num += 1
    row_idx = num
    ws.cell(row=row_idx, column=1, value=num - 1).border = thin_border
    ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center")
    ws.cell(row=row_idx, column=2, value=tarih).border = thin_border
    ws.cell(row=row_idx, column=3, value=firma).border = thin_border
    ws.cell(row=row_idx, column=4, value=pos).border = thin_border
    ec = ws.cell(row=row_idx, column=5, value=ergebnis)
    ec.border = thin_border
    ec.font = Font(name="Calibri", bold=True, size=10)
    f = ergebnis_fill(ergebnis)
    if f: ec.fill = f
    ws.cell(row=row_idx, column=6, value=notiz).border = thin_border

# Autofilter güncelle
ws.auto_filter.ref = f"A1:F{num}"

wb.save(EXCEL)
print(f"✅ {len(new_entries)} yeni satır eklendi. Toplam: {num - 1}")
