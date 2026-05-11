import openpyxl
from copy import copy

wb = openpyxl.load_workbook(r'C:\Users\hamza\Desktop\Lebenslauf\Bewerbungen_Uebersicht.xlsx')
ws = wb.active

# Row numbers (column A "#") to remove — non-application entries
remove_ids = {
    # WEB.DE newsletters / marketing
    1, 6, 23, 36, 67, 91, 112, 161,
    # Google account
    17, 21, 39,
    # Microsoft account / OneDrive
    27, 28, 29, 32, 40, 99,
    # eSIM.sm (purchases, activation, marketing)
    68, 69, 70, 71, 72, 136,
    # DHL delivery
    80,
    # Amazon (orders, shipping, Prime, Luna, marketing)
    131, 132, 134, 135, 144, 145, 147, 148, 162,
    # Procter & Gamble refunds (Ariel / Lenor)
    115, 116,
    # GitGuardian / GitHub security
    130, 137, 138, 139, 140, 141, 142, 160,
    # Cloudflare DNS
    146,
    # Netlify
    133,
    # GermanTechJobs newsletter
    48, 110,
    # Instaffo account management (not application)
    63,
    # Instaffo pure job suggestions (no actual application)
    81, 85, 127, 143, 158,
    # Dicom / former employer documents (not a Bewerbung)
    126,
    # Bundesagentur für Arbeit Jobalert (notification, not application)
    157,
}

print(f"Einträge zum Entfernen: {len(remove_ids)}")

# Collect rows to keep (data rows start at row 2)
rows_to_keep = []
rows_to_remove = []
for row in range(2, ws.max_row + 1):
    nr = ws.cell(row=row, column=1).value
    if nr is None:
        continue
    try:
        nr_int = int(nr)
    except (ValueError, TypeError):
        rows_to_keep.append(row)
        continue
    if nr_int in remove_ids:
        firma = ws.cell(row=row, column=3).value or ""
        pos = ws.cell(row=row, column=4).value or ""
        rows_to_remove.append((nr_int, firma, pos))
    else:
        rows_to_keep.append(row)

print(f"\nEntfernte Einträge ({len(rows_to_remove)}):")
for nr, firma, pos in sorted(rows_to_remove):
    print(f"  #{nr}: {firma} | {pos}")

print(f"\nVerbleibende Einträge: {len(rows_to_keep)}")

# Build new workbook with same structure
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = ws.title

# Copy header row with formatting
for col in range(1, ws.max_column + 1):
    src = ws.cell(row=1, column=col)
    dst = ws2.cell(row=1, column=col, value=src.value)
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.alignment = copy(src.alignment)
    dst.border = copy(src.border)

# Copy column widths
for col_letter, dim in ws.column_dimensions.items():
    ws2.column_dimensions[col_letter].width = dim.width

# Copy kept rows with new numbering
new_nr = 1
for old_row in rows_to_keep:
    new_row = new_nr + 1  # +1 for header
    for col in range(1, ws.max_column + 1):
        src = ws.cell(row=old_row, column=col)
        if col == 1:
            dst = ws2.cell(row=new_row, column=col, value=new_nr)
        else:
            dst = ws2.cell(row=new_row, column=col, value=src.value)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.alignment = copy(src.alignment)
        dst.border = copy(src.border)
    new_nr += 1

# Freeze top row + auto-filter
ws2.freeze_panes = 'A2'
ws2.auto_filter.ref = f"A1:F{ws2.max_row}"

wb2.save(r'C:\Users\hamza\Desktop\Lebenslauf\Bewerbungen_Uebersicht.xlsx')
print(f"\nGespeichert! {new_nr - 1} Einträge in Bewerbungen_Uebersicht.xlsx")
